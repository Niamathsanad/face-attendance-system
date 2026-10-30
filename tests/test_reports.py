"""Report generation, CSV/Excel export, timezone handling, export safety."""
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import new_student
from services import reporting_service as rs
from services import time_utils as tu
from services.time_utils import now_utc_iso
from services.validators import ValidationError


@pytest.fixture
def populated(db):
    a = new_student(db, "A1", "Ali", "CS", "Sem-5"); b = new_student(db, "B1", "=HYPERLINK(\"http://x\")", "EE", "Sem-3")
    db.execute("UPDATE students SET registration_date='2026-01-01'"); db.commit()

    def sess(name, day):
        c = db.execute("INSERT INTO attendance_sessions (name, session_date, status, started_at) VALUES (?,?, 'closed', ?)", (name, day, now_utc_iso()))
        db.commit(); return c.lastrowid
    s1, s2, s3 = sess("Mon", "2026-02-02"), sess("Tue", "2026-02-03"), sess("Next month", "2026-03-04")
    for s, sid in [(s1, a), (s1, b), (s2, a), (s3, b)]:
        db.execute("INSERT INTO attendance_records (session_id, student_id, status, source, marked_at) VALUES (?,?, 'Present','face','2026-02-02T04:00:00+00:00')", (s, sid))
    db.commit()
    return db


def test_day_report(populated):
    r = rs.build_report(populated, "day", "2026-02-02")
    assert r["summary"] == {"expected": 2, "present": 2, "absent": 0, "percentage": 100.0}


def test_week_report_covers_monday_to_sunday(populated):
    r = rs.build_report(populated, "week", "2026-02-04")           # a Wednesday
    assert (r["date_from"], r["date_to"]) == ("2026-02-02", "2026-02-08")
    assert r["summary"]["expected"] == 4 and r["summary"]["present"] == 3


def test_month_report(populated):
    assert rs.build_report(populated, "month", "2026-02-15")["summary"]["expected"] == 4
    assert rs.build_report(populated, "month", "2026-03-15")["summary"]["present"] == 1


def test_student_report(populated):
    r = rs.build_report(populated, "student", student_code="A1", date_from="2026-02-01", date_to="2026-03-31")
    assert r["summary"]["expected"] == 3 and r["summary"]["present"] == 2          # absent in the March session
    with pytest.raises(ValidationError):
        rs.build_report(populated, "student", student_code="NOPE")
    with pytest.raises(ValidationError):
        rs.build_report(populated, "student")


def test_group_report(populated):
    r = rs.build_report(populated, "group", department="EE", date_from="2026-02-01", date_to="2026-03-31")
    assert set(r["frame"]["student_code"]) == {"B1"}
    with pytest.raises(ValidationError):
        rs.build_report(populated, "group")


def test_csv_export_content_and_timezone(populated):
    r = rs.build_report(populated, "day", "2026-02-02")
    csv = rs.export_csv(r).decode("utf-8-sig")
    assert "Marked At (Asia/Karachi)" in csv.splitlines()[0]
    assert "2026-02-02 09:00:00" in csv                  # 04:00 UTC == 09:00 in Pakistan (UTC+5)


def test_excel_export_has_three_sheets(populated):
    r = rs.build_report(populated, "month", "2026-02-15")
    wb = load_workbook(io.BytesIO(rs.export_excel(r)))
    assert wb.sheetnames == ["Summary", "Attendance", "Per student"]
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Timezone"] == "Asia/Karachi" and summary["Present"] == 3


def test_formula_injection_is_neutralised_in_exports(populated):
    r = rs.build_report(populated, "month", "2026-02-15")
    csv = rs.export_csv(r).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv and ",=HYPERLINK" not in csv
    wb = load_workbook(io.BytesIO(rs.export_excel(r)))
    names = [c.value for c in wb["Attendance"]["D"]]
    assert all(not str(n).startswith("=") for n in names if n)


def test_export_routes(auth_client, populated):
    r = auth_client.get("/reports/export?kind=day&ref_date=2026-02-02&format=csv")
    assert r.status_code == 200 and r.mimetype == "text/csv" and "attachment" in r.headers["Content-Disposition"]
    r = auth_client.get("/reports/export?kind=day&ref_date=2026-02-02&format=xlsx")
    assert r.status_code == 200 and r.data[:2] == b"PK"            # xlsx files are zip containers
    assert auth_client.get("/reports/export?kind=student&student_code=NOPE").status_code == 400


def test_reports_page_preview(auth_client, populated):
    html = auth_client.get("/reports?run=1&kind=week&ref_date=2026-02-04").get_data(as_text=True)
    assert "Weekly report" in html and "2026-02-02 to 2026-02-08" in html


def test_timezone_conversion_and_session_date():
    assert tu.to_local_str("2026-01-01T00:00:00+00:00") == "2026-01-01 05:00:00"
    assert tu.to_local_str(None) == ""
    assert tu.tz_name() == "Asia/Karachi"
    assert tu.week_bounds(tu.parse_date("2026-02-04"))[0].isoformat() == "2026-02-02"
    assert tu.month_bounds(tu.parse_date("2026-02-15"))[1].isoformat() == "2026-02-28"
